# from docx import Document
# from PIL import Image
# import pytesseract
# import os
# from pypdf import PdfReader


# # -------------------------
# # DOCX
# # -------------------------
# def extract_docx(path):
#     doc = Document(path)
#     return "\n".join([p.text for p in doc.paragraphs])


# # -------------------------
# # IMAGE OCR
# # -------------------------
# def extract_image(path):
#     img = Image.open(path)
#     return pytesseract.image_to_string(img)


# # -------------------------
# # PDF TEXT
# # -------------------------
# def extract_pdf(path):
#     reader = PdfReader(path)
#     text = ""

#     for page in reader.pages:
#         page_text = page.extract_text()
#         if page_text:
#             text += page_text + "\n"

#     return text


# # -------------------------
# # ROUTER
# # -------------------------
# def extract_text(path: str) -> str:
#     ext = os.path.splitext(path)[1].lower()

#     if ext == ".docx":
#         return extract_docx(path)

#     elif ext in [".png", ".jpg", ".jpeg"]:
#         return extract_image(path)

#     elif ext == ".pdf":
#         return extract_pdf(path)

#     elif ext == ".txt":
#         with open(path, "r", encoding="utf-8") as f:
#             return f.read()

#     else:
#         raise ValueError(f"Unsupported file type: {ext}")


# # -------------------------
# # INGESTION NODE
# # -------------------------
# def ingestion_node(state):
#     file_path = state["file_path"]

#     raw_text = extract_text(file_path)

#     return {
#         **state,
#         "raw_text": raw_text
#     }


import io
from pypdf import PdfReader
from docx import Document
from PIL import Image
import pytesseract
from graph.state import MedifyState

def ingestion_node(state: MedifyState):
    """
    Extracts raw text from uploaded files stored in memory.
    Supports: .pdf, .docx, .png, .jpg, .jpeg, .txt
    """
    # 1. Retrieve the bytes and filename from the graph state
    file_bytes = state.get("file_bytes")
    file_name = state.get("file_name", "").lower()
    
    if not file_bytes:
        return {"raw_text": "Error: No file content provided."}

    text = ""

    # 2. Wrap bytes in a BytesIO stream to simulate a file object
    # This allows libraries like pypdf and python-docx to read the data without a disk path
    file_stream = io.BytesIO(file_bytes)

    try:
        # --- PDF Processing ---
        if file_name.endswith(".pdf"):
            reader = PdfReader(file_stream)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        
        # --- Word Document Processing ---
        elif file_name.endswith(".docx"):
            doc = Document(file_stream)
            text = "\n".join([para.text for para in doc.paragraphs])
        
        # --- Image Processing (OCR) ---
        elif file_name.endswith((".png", ".jpg", ".jpeg")):
            image = Image.open(file_stream)
            text = pytesseract.image_to_string(image)
        
        # --- Plain Text Processing ---
        elif file_name.endswith(".txt"):
            # Bytes are decoded directly into a string
            text = file_bytes.decode("utf-8")
        
        # --- Unsupported Types ---
        else:
            text = f"Unsupported file format: {file_name}. Please upload a PDF, DOCX, TXT, or Image."

    except Exception as e:
        # We catch the exception so the LangGraph doesn't crash entirely
        # This allows the 'gate' node to handle the failure gracefully
        text = f"Error during text extraction: {str(e)}"

    # 3. Return the extracted text to be stored in the state
    # In LangGraph, returning a dict updates the shared state
    return {"raw_text": text}