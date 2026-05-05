# import fitz
# from docx import Document
# from PIL import Image
# import pytesseract


# def process_document(file_path: str) -> str:

#     if file_path.endswith(".pdf"):
#         doc = fitz.open(file_path)
#         text = ""
#         for page in doc:
#             text += page.get_text()
#         return text

#     elif file_path.endswith(".docx"):
#         doc = Document(file_path)
#         return " ".join([p.text for p in doc.paragraphs])

#     elif file_path.endswith((".png", ".jpg", ".jpeg")):
#         img = Image.open(file_path)
#         return pytesseract.image_to_string(img)

#     return ""

import fitz
from docx import Document
from PIL import Image
import pytesseract
import io


def process_document(file_bytes: bytes, filename: str) -> str:

    try:
        if filename.endswith(".pdf"):
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            text = ""
            for page in doc:
                text += page.get_text()
            return text

        elif filename.endswith(".docx"):
            doc = Document(io.BytesIO(file_bytes))
            return " ".join([p.text for p in doc.paragraphs])

        elif filename.endswith((".png", ".jpg", ".jpeg")):
            img = Image.open(io.BytesIO(file_bytes))
            return pytesseract.image_to_string(img)

        return ""

    except Exception as e:
        print("🔥 Processing error:", str(e))
        return ""